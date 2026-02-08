# SIMPLIFIED VERSION - No FFmpeg required
import sounddevice as sd
import numpy as np
import wave
import speech_recognition as sr
from tkinter import *
from tkinter import ttk, messagebox, filedialog, scrolledtext
from docx import Document
from datetime import datetime
import threading
import tempfile
import os
import queue
import time
import json
import re

# Global variables
is_recording = False
is_live_transcribing = False
audio_queue = queue.Queue()
fs = 44100
channels = 1
chunk_size = 1024

class AudioProcessor:
    @staticmethod
    def normalize_volume(audio_data):
        """Normalize audio volume"""
        max_val = np.max(np.abs(audio_data))
        if max_val > 0:
            return (audio_data / max_val * 32767).astype(np.int16)
        return audio_data

    @staticmethod
    def trim_silence(audio_data, threshold=0.01):
        """Trim silence from beginning and end"""
        audio_flat = audio_data.flatten().astype(np.float32) / 32768.0
        above_threshold = np.where(np.abs(audio_flat) > threshold)[0]
        if len(above_threshold) > 0:
            start = max(0, above_threshold[0] - 1000)
            end = min(len(audio_flat), above_threshold[-1] + 1000)
            return audio_data[start:end]
        return audio_data

class GrammarCorrector:
    def __init__(self):
        self.business_terms = {
            'kpi': 'KPI', 'roi': 'ROI', 'asap': 'ASAP', 'eta': 'ETA',
            'fyi': 'for your information', 'btw': 'by the way',
            'imho': 'in my humble opinion', 'tbd': 'to be determined'
        }
        
    def correct_grammar(self, text):
        """Basic grammar correction"""
        sentences = text.split('. ')
        corrected_sentences = []
        
        for sentence in sentences:
            if sentence.strip():
                sentence = sentence.strip()
                if sentence:
                    sentence = sentence[0].upper() + sentence[1:]
                if not sentence.endswith(('.', '!', '?')):
                    sentence += '.'
                corrected_sentences.append(sentence)
        
        corrected_text = '. '.join(corrected_sentences)
        corrected_text = self.apply_business_terms(corrected_text)
        return corrected_text
    
    def apply_business_terms(self, text):
        for informal, formal in self.business_terms.items():
            text = re.sub(r'\b' + re.escape(informal) + r'\b', formal, text, flags=re.IGNORECASE)
        return text

class EnterpriseVoiceRecognitionSystem:
    def __init__(self, root):
        self.root = root
        self.setup_enterprise_features()
        self.setup_gui()
        self.audio_processor = AudioProcessor()
        self.grammar_corrector = GrammarCorrector()
        self.current_audio_data = None
        
    def setup_enterprise_features(self):
        self.projects = {}
        self.team_members = {}
        self.session_history = []
        
    def setup_gui(self):
        self.root.title("🏢 Enterprise Voice Recognition Suite")
        self.root.geometry("1000x700")
        self.root.configure(bg='#1e2a36')
        
        self.notebook = ttk.Notebook(self.root)
        self.notebook.pack(fill=BOTH, expand=True, padx=10, pady=10)
        
        self.recording_tab = Frame(self.notebook, bg='#1e2a36')
        self.grammar_tab = Frame(self.notebook, bg='#1e2a36')
        self.enterprise_tab = Frame(self.notebook, bg='#1e2a36')
        
        self.notebook.add(self.recording_tab, text="🎙️ Recording")
        self.notebook.add(self.grammar_tab, text="✍️ Grammar & Business")
        self.notebook.add(self.enterprise_tab, text="🏢 Enterprise")
        
        self.setup_recording_tab()
        self.setup_grammar_tab()
        self.setup_enterprise_tab()

    def setup_recording_tab(self):
        header = Frame(self.recording_tab, bg='#2c3e50')
        header.pack(fill=X, padx=10, pady=10)
        
        Label(header, text="Professional Voice Recognition", 
              font=("Arial", 16, "bold"), bg="#2c3e50", fg="white").pack(pady=10)
        
        control_panel = Frame(self.recording_tab, bg='#1e2a36')
        control_panel.pack(fill=X, padx=20, pady=10)
        
        # Language selection
        lang_frame = Frame(control_panel, bg='#1e2a36')
        lang_frame.pack(pady=5)
        
        Label(lang_frame, text="Language:", font=("Arial", 10, "bold"), 
              bg="#1e2a36", fg="white").pack(side=LEFT, padx=5)
        
        self.language_var = StringVar(value="English")
        languages = ["English", "Urdu", "Turkish", "Arabic", "Spanish", "French"]
        
        self.language_menu = ttk.Combobox(lang_frame, textvariable=self.language_var, 
                                        values=languages, state="readonly", width=15)
        self.language_menu.pack(side=LEFT, padx=5)
        
        # Buttons
        btn_frame = Frame(control_panel, bg='#1e2a36')
        btn_frame.pack(pady=15)
        
        self.start_btn = Button(btn_frame, text="🎙️ Start Recording", 
                               command=self.start_recording, bg="#27ae60", fg="white",
                               font=("Arial", 10, "bold"), width=15, height=1)
        self.start_btn.pack(side=LEFT, padx=5)
        
        self.stop_btn = Button(btn_frame, text="🛑 Stop Recording", 
                              command=self.stop_recording, bg="#e74c3c", fg="white",
                              font=("Arial", 10, "bold"), width=15, height=1, state=DISABLED)
        self.stop_btn.pack(side=LEFT, padx=5)
        
        Button(btn_frame, text="📁 Import WAV", command=self.import_audio_file, 
               bg="#3498db", fg="white", font=("Arial", 10, "bold"), width=15, height=1).pack(side=LEFT, padx=5)
        
        # Status
        self.status_label = Label(control_panel, text="Ready to record", 
                                 font=("Arial", 11, "bold"), bg="#1e2a36", fg="#2ecc71")
        self.status_label.pack(pady=5)
        
        # Text area
        text_frame = Frame(self.recording_tab, bg='#1e2a36')
        text_frame.pack(fill=BOTH, expand=True, padx=20, pady=10)
        
        self.text_area = scrolledtext.ScrolledText(text_frame, wrap=WORD, 
                                                  font=("Arial", 11), 
                                                  bg="white", fg="#2c3e50",
                                                  height=15)
        self.text_area.pack(fill=BOTH, expand=True)

    def setup_grammar_tab(self):
        Label(self.grammar_tab, text="Grammar & Business Tools", 
              font=("Arial", 16, "bold"), bg="#1e2a36", fg="white").pack(pady=15)
        
        grammar_frame = Frame(self.grammar_tab, bg='#1e2a36')
        grammar_frame.pack(fill=X, padx=20, pady=10)
        
        Button(grammar_frame, text="✅ Auto-Correct Grammar", 
               command=self.auto_correct_grammar, bg="#9b59b6", fg="white",
               font=("Arial", 10, "bold"), width=20).pack(side=LEFT, padx=5)
        
        Button(grammar_frame, text="💼 Business Format", 
               command=self.apply_business_format, bg="#1abc9c", fg="white",
               font=("Arial", 10, "bold"), width=15).pack(side=LEFT, padx=5)
        
        self.grammar_text = scrolledtext.ScrolledText(self.grammar_tab, wrap=WORD, 
                                                     font=("Arial", 11), 
                                                     bg="white", fg="#2c3e50",
                                                     height=15)
        self.grammar_text.pack(fill=BOTH, expand=True, padx=20, pady=10)

    def setup_enterprise_tab(self):
        Label(self.enterprise_tab, text="Enterprise Dashboard", 
              font=("Arial", 16, "bold"), bg="#1e2a36", fg="white").pack(pady=15)
        
        btn_frame = Frame(self.enterprise_tab, bg='#1e2a36')
        btn_frame.pack(fill=X, padx=20, pady=10)
        
        Button(btn_frame, text="📊 Usage Analytics", 
               command=self.show_analytics, bg="#3498db", fg="white",
               font=("Arial", 10, "bold"), width=15).pack(side=LEFT, padx=5)
        
        Button(btn_frame, text="💾 Export Document", 
               command=self.export_document, bg="#27ae60", fg="white",
               font=("Arial", 10, "bold"), width=15).pack(side=LEFT, padx=5)
        
        self.dashboard_text = scrolledtext.ScrolledText(self.enterprise_tab, wrap=WORD, 
                                                       font=("Arial", 10), 
                                                       bg="white", fg="#2c3e50",
                                                       height=15)
        self.dashboard_text.pack(fill=BOTH, expand=True, padx=20, pady=10)
        self.update_dashboard()

    def audio_callback(self, indata, frames, time, status):
        if is_recording:
            audio_queue.put(indata.copy())

    def start_recording(self):
        global is_recording
        is_recording = True
        self.status_label.config(text="🔴 RECORDING - Speak now...", fg="#e74c3c")
        self.start_btn.config(state=DISABLED)
        self.stop_btn.config(state=NORMAL)
        
        while not audio_queue.empty():
            audio_queue.get()
        
        threading.Thread(target=self.record_audio, daemon=True).start()

    def stop_recording(self):
        global is_recording
        is_recording = False
        self.status_label.config(text="⏳ Processing audio...", fg="#f39c12")
        threading.Thread(target=self.process_recorded_audio, daemon=True).start()

    def record_audio(self):
        try:
            with sd.InputStream(samplerate=fs, channels=channels, 
                              dtype='int16', callback=self.audio_callback,
                              blocksize=chunk_size):
                while is_recording:
                    time.sleep(0.1)
        except Exception as e:
            self.root.after(0, lambda: messagebox.showerror("Error", f"Recording failed: {e}"))

    def process_recorded_audio(self):
        audio_chunks = []
        while not audio_queue.empty():
            chunk = audio_queue.get()
            audio_chunks.append(chunk)
        
        if not audio_chunks:
            self.root.after(0, lambda: messagebox.showwarning("Warning", "No audio recorded!"))
            self.reset_recording_buttons()
            return
        
        try:
            audio_data = np.concatenate(audio_chunks, axis=0)
            self.current_audio_data = audio_data
            
            # Enhance audio
            audio_data = self.audio_processor.normalize_volume(audio_data)
            audio_data = self.audio_processor.trim_silence(audio_data)
            
            # Save temp file
            temp_file = tempfile.mktemp(suffix=".wav")
            with wave.open(temp_file, 'wb') as wf:
                wf.setnchannels(channels)
                wf.setsampwidth(2)
                wf.setframerate(fs)
                wf.writeframes(audio_data.tobytes())
            
            self.transcribe_audio(temp_file)
            self.session_history.append(f"Recording {datetime.now().strftime('%H:%M:%S')}")
            self.update_dashboard()
            
        except Exception as e:
            self.root.after(0, lambda: messagebox.showerror("Error", f"Processing failed: {e}"))
        finally:
            self.reset_recording_buttons()

    def transcribe_audio(self, audio_file):
        try:
            recognizer = sr.Recognizer()
            with sr.AudioFile(audio_file) as source:
                recognizer.adjust_for_ambient_noise(source, duration=0.5)
                audio = recognizer.record(source)
            
            lang_code = self.get_language_code()
            text = recognizer.recognize_google(audio, language=lang_code)
            
            self.root.after(0, lambda: self.text_area.insert(END, text + "\n\n"))
            self.root.after(0, lambda: self.grammar_text.insert(END, text + "\n\n"))
            self.root.after(0, lambda: self.status_label.config(text="✅ Transcription Complete!", fg="#2ecc71"))
            
            if os.path.exists(audio_file):
                os.remove(audio_file)
                
        except sr.UnknownValueError:
            self.root.after(0, lambda: messagebox.showerror("Error", "Could not understand audio"))
        except Exception as e:
            self.root.after(0, lambda: messagebox.showerror("Error", f"Transcription failed: {e}"))

    def import_audio_file(self):
        file_path = filedialog.askopenfilename(filetypes=[("WAV Files", "*.wav")])
        if file_path and file_path.endswith('.wav'):
            try:
                self.status_label.config(text="⏳ Processing WAV file...", fg="#f39c12")
                threading.Thread(target=self.process_wav_file, args=(file_path,), daemon=True).start()
            except Exception as e:
                messagebox.showerror("Error", f"Import failed: {e}")

    def process_wav_file(self, file_path):
        try:
            with wave.open(file_path, 'rb') as wf:
                frames = wf.getnframes()
                audio_data = np.frombuffer(wf.readframes(frames), dtype=np.int16)
                self.current_audio_data = audio_data.reshape(-1, 1)
            
            temp_file = tempfile.mktemp(suffix=".wav")
            with wave.open(temp_file, 'wb') as wf:
                wf.setnchannels(1)
                wf.setsampwidth(2)
                wf.setframerate(16000)
                wf.writeframes(audio_data.tobytes())
            
            self.transcribe_audio(temp_file)
        except Exception as e:
            self.root.after(0, lambda: messagebox.showerror("Error", f"WAV processing failed: {e}"))

    def auto_correct_grammar(self):
        text = self.grammar_text.get(1.0, END).strip()
        if text:
            corrected = self.grammar_corrector.correct_grammar(text)
            self.grammar_text.delete(1.0, END)
            self.grammar_text.insert(END, corrected)
            messagebox.showinfo("Success", "Grammar corrected!")

    def apply_business_format(self):
        text = self.grammar_text.get(1.0, END).strip()
        if text:
            sentences = text.split('. ')
            formatted = []
            for sentence in sentences:
                if sentence.strip():
                    s = sentence.strip()
                    s = s[0].upper() + s[1:]
                    if not s.endswith('.'): s += '.'
                    formatted.append(s)
            self.grammar_text.delete(1.0, END)
            self.grammar_text.insert(END, '. '.join(formatted))
            messagebox.showinfo("Success", "Business formatting applied!")

    def show_analytics(self):
        analytics = f"📊 USAGE ANALYTICS\n\n"
        analytics += f"Total Recordings: {len(self.session_history)}\n"
        analytics += f"Projects: {len(self.projects)}\n"
        analytics += f"Team Members: {len(self.team_members)}\n"
        analytics += f"System: ✅ Operational\n"
        self.dashboard_text.delete(1.0, END)
        self.dashboard_text.insert(END, analytics)

    def export_document(self):
        text = self.grammar_text.get(1.0, END).strip()
        if text:
            file_path = filedialog.asksaveasfilename(
                defaultextension=".docx",
                filetypes=[("Word Document", "*.docx")],
                initialfile=f"transcript_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
            )
            if file_path:
                try:
                    doc = Document()
                    doc.add_paragraph(text)
                    doc.save(file_path)
                    messagebox.showinfo("Success", f"Exported to:\n{file_path}")
                except Exception as e:
                    messagebox.showerror("Error", f"Export failed: {e}")

    def get_language_code(self):
        lang_map = {
            "English": "en-US", "Urdu": "ur-PK", "Turkish": "tr-TR",
            "Arabic": "ar-SA", "Spanish": "es-ES", "French": "fr-FR"
        }
        return lang_map.get(self.language_var.get(), "en-US")

    def reset_recording_buttons(self):
        self.start_btn.config(state=NORMAL)
        self.stop_btn.config(state=DISABLED)
        self.status_label.config(text="Ready to record", fg="#2ecc71")

    def update_dashboard(self):
        info = f"🏢 ENTERPRISE DASHBOARD\n\n"
        info += f"Recordings: {len(self.session_history)}\n"
        info += f"Projects: {len(self.projects)}\n"
        info += f"Status: ✅ System Ready\n"
        info += f"Last Update: {datetime.now().strftime('%H:%M:%S')}\n"
        self.dashboard_text.delete(1.0, END)
        self.dashboard_text.insert(END, info)

# Run the application
if __name__ == "__main__":
    try:
        # Test microphone
        with sd.InputStream(samplerate=44100): pass
        print("Microphone access confirmed!")
    except:
        messagebox.showwarning("Microphone", "Check microphone permissions")
    
    root = Tk()
    app = EnterpriseVoiceRecognitionSystem(root)
    root.mainloop()